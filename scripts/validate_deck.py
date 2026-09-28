#!/usr/bin/env python3
"""
validate_deck.py — structural QA for a First Chair deck.

    python3 scripts/validate_deck.py Deck.pptx

Checks, in order of how often each one has actually bitten:

  1. ZIP integrity            — the file opens as an archive at all.
  2. python-pptx parse        — the OOXML is well-formed enough to load.
  3. Canvas size              — must be 13.333 x 7.5 in (LAYOUT_WIDE).
  4. Off-canvas shapes        — pptxgenjs writes out-of-bounds coords silently.
  5. Font drift               — every run should be Georgia (headlines, figures)
                                or Calibri (everything else).
  6. Palette drift            — flags colors outside the house token set.
  7. Footer / page numbers    — content slides should carry both (any text; the
                                footer string is per-deck and comes from the brief).

Exit code 0 = clean, 1 = at least one FAIL. WARN never fails the build.
This does NOT catch text overflow or overlap. Render to images and look.
"""

import sys
import zipfile
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:
    sys.exit("python-pptx is required:  pip install python-pptx")

PALETTE = {
    # mirrors assets/deck_style_kit.js PALETTE (v2); regenerate with:
    # node -e "console.log(require('./assets/deck_style_kit.js').PALETTE.join(' '))"
    "F6F7F5", "FFFFFF", "0E2A47", "081A2D", "0D9488", "0F766E", "5EEAD4", "E3EFEC",
    "F59E0B", "B45309", "16232E", "5A6B7B", "C9D6E2", "9FB3C4", "D5DBE0", "ECEFF1",
    "15803D", "DCFCE7", "14532D", "D97706", "FEF3C7", "78350F", "B91C1C", "FEE2E2",
    "7F1D1D",
    "000000",
}
FONTS = {"Georgia", "Calibri"}
EMU_IN = 914400


def true_box(sh):
    """(left, top, right, bottom) in inches as drawn. A box rotated a quarter
    turn occupies its height across and its width down, about the same centre
    (the matrix y-axis label is drawn that way)."""
    l, t, w, h = sh.left / EMU_IN, sh.top / EMU_IN, (sh.width or 0) / EMU_IN, (sh.height or 0) / EMU_IN
    rot = round((getattr(sh, "rotation", 0) or 0) % 180)
    if rot == 90:
        cx, cy = l + w / 2, t + h / 2
        l, t, w, h = cx - h / 2, cy - w / 2, h, w
    return (l, t, l + w, t + h)
FAILS, WARNS = [], []


def fail(msg):
    FAILS.append(msg)


def warn(msg):
    WARNS.append(msg)


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6 and hasattr(sh, "shapes"):  # group
            yield from walk(sh.shapes)


def main(path):
    p = Path(path)
    if not p.exists():
        sys.exit(f"no such file: {p}")

    # 1. zip integrity
    try:
        bad = zipfile.ZipFile(p).testzip()
        if bad:
            fail(f"corrupt archive member: {bad}")
    except zipfile.BadZipFile:
        sys.exit("FAIL: not a valid .pptx archive (did something rezip it?)")

    # 2. parse
    try:
        prs = Presentation(str(p))
    except Exception as exc:  # noqa: BLE001
        sys.exit(f"FAIL: python-pptx cannot open the deck: {exc}")

    # 2b. render stamp — a deck without it was not built by render_deck.js
    subj = prs.core_properties.subject or ""
    if "firstchair-consulting-deck render" not in subj:
        fail("no render stamp in the file's subject field: this deck was not produced by scripts/render_deck.js. "
             "Hand-rolled pptxgenjs code is not a deliverable of this skill; build from a brief.")
    # 2c. speaker notes on every content slide (the candid channel; required)
    for n, slide in enumerate(prs.slides, 1):
        texts = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame).strip()
        structural = n == 1 or len(texts) < 60
        has_notes = slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()
        if not structural and not has_notes:
            fail(f"slide {n}: no speaker notes; every content slide carries them (what the slide says diplomatically, the notes say plainly)")

    # 3. canvas
    w_in = prs.slide_width / EMU_IN
    h_in = prs.slide_height / EMU_IN
    if not (13.2 < w_in < 13.4 and 7.4 < h_in < 7.6):
        fail(f"canvas is {w_in:.2f}x{h_in:.2f} in — set pres.layout='LAYOUT_WIDE' BEFORE adding slides")

    for i, slide in enumerate(prs.slides, start=1):
        has_footer = has_page = False
        for sh in walk(slide.shapes):
            # 4. off-canvas
            try:
                if sh.left is None or sh.top is None:
                    continue
                l, t, r, b = true_box(sh)
                if l < -0.01 or t < -0.01 or r > w_in + 0.02 or b > h_in + 0.02:
                    fail(f"slide {i}: shape '{sh.shape_id}' extends off-canvas "
                         f"(right {r:.2f}\", bottom {b:.2f}\")")
            except (TypeError, ValueError):
                pass

            if not sh.has_text_frame:
                continue
            text = sh.text_frame.text.strip()
            in_footer_band = sh.top is not None and 6.85 < (sh.top / EMU_IN) < 7.35
            if in_footer_band and text and not text.isdigit():
                has_footer = True
            if in_footer_band and text.isdigit():
                has_page = True

            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    # 5. font
                    name = run.font.name
                    if name and name not in FONTS:
                        warn(f"slide {i}: off-style font ({name}); the house style is Georgia and Calibri: {run.text[:32]!r}")
                    # 6. palette
                    try:
                        rgb = run.font.color.rgb
                    except (AttributeError, TypeError):
                        rgb = None
                    if rgb is not None and str(rgb).upper() not in PALETTE:
                        warn(f"slide {i}: off-palette text color {rgb} on {run.text[:32]!r}")

        # 7. footer (title slide is exempt)
        if i > 1 and not (has_footer and has_page):
            warn(f"slide {i}: missing footer text and/or page number")

    print(f"{p.name}: {len(prs.slides)} slides, {w_in:.2f}x{h_in:.2f} in")
    for w in WARNS:
        print("  WARN", w)
    for f in FAILS:
        print("  FAIL", f)
    if FAILS:
        print("\nFAILED — fix in the generator script and rebuild. Never hand-edit the packed XML.")
        return 1
    print("\nPASS (structure only) — now render to images and check overflow, overlap, alignment.")
    return 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    sys.exit(main(sys.argv[1]))
